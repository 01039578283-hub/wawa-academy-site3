"""Recheck authorized source hashes, mappings and field-level differences."""
from pathlib import Path
import csv,hashlib,json,re
from openpyxl import load_workbook
import improve_neighborhood_pages as impl

def main():
 source=Path(r'C:\Users\1992k\Desktop\홈페이지 작업 폴더\센터정보');files=[]
 for record in impl.FACTS['sources']:
  p=source/record['file'];sha=hashlib.sha256(p.read_bytes()).hexdigest();assert sha==record['sha256'],p
  files.append({**record,'matchesIntegratedSnapshot':True})
 wb=load_workbook(source/'코칭센터_데이터_.xlsx',read_only=True,data_only=True)
 rows=list(wb.active.values);by_name={str(r[0]).strip():(i,r) for i,r in enumerate(rows,1) if r[0]};wb.close()
 differences=[];unmatched=[];identity=[]
 for c in impl.FACTS['centers']:
  match=by_name.get(c['sourceName'])
  if not match:unmatched.append({'center':c['routeName'],'sourceName':c['sourceName'],'behavior':'Use CSV explicitly; confirm current course conditions'});continue
  i,row=match
  for field,col in [('address',11),('registeredName',6),('registrationNumber',7)]:
   if impl.norm(c[field])!=impl.norm(str(row[col] or '')):identity.append({'center':c['routeName'],'field':field})
  for subject,col in zip(['국어','영어','수학','과학','사회'],range(16,21)):
   actual=re.findall(r'[초중고][1-6]',str(row[col] or ''))
   csv_subjects=c.get('csvSubjects',c['subjects'])
   if csv_subjects[subject]!=actual:differences.append({'center':c['routeName'],'region':c['region'],'subject':subject,'csv':csv_subjects[subject],'workbook':actual,'workbookRow':i})
 assert not identity,identity
 with (source/'센터 정보 및 교육비 371개 코드_최신화.csv').open(encoding='utf-8-sig',newline='') as f:snippets=list(csv.reader(f))
 with (source/'센터정보 정리.csv').open(encoding='utf-8-sig',newline='') as f:areas=list(csv.reader(f))[1:]
 schools_wb=load_workbook(source/'타깃학교 v2.xlsx',read_only=True,data_only=True)
 schools=list(schools_wb.active.values)[1:];schools_wb.close()
 assert len(snippets)==len(areas)==len(schools)==371
 mapping=[]
 for i,(area,school) in enumerate(zip(areas,schools),1):
  assert impl.norm(area[0])==impl.norm(school[0]),('area',i)
  assert impl.norm(area[6])==impl.norm(school[3]),('center',i)
  mapping.append({'row':i,'neighborhood':area[0],'sourceCenter':area[6]})
 impl.dump(impl.DATA/'source-conflicts.json',differences)
 impl.dump(impl.REPORT/'source-reconciliation.json',{'sourceFiles':files,'sourceNeighborhoodRows':371,'integratedCenters':len(impl.FACTS['centers']),'matchedWorkbookCenters':len(impl.FACTS['centers'])-len(unmatched),'differentGradeFields':len(differences),'differentGradeCenters':len({v['center'] for v in differences}),'identityConflicts':identity,'unmatchedWorkbookCenters':unmatched,'neighborhoodMapping':mapping})
 print(json.dumps({'sourceHashesVerified':len(files),'neighborhoodRowsVerified':371,'differentGradeFields':len(differences),'differentGradeCenters':len({v['center'] for v in differences}),'identityConflicts':identity,'unmatchedWorkbookCenters':unmatched},ensure_ascii=False))
if __name__=='__main__':main()
